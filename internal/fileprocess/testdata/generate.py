"""Generate synthetic application-library fixtures. Refuse to overwrite any original.
Run with the bundled authoring Python (python-docx, python-pptx, openpyxl).
These are test data, not runtime dependencies.
"""
from pathlib import Path
from datetime import datetime, timezone
from zipfile import ZipFile, ZIP_STORED, ZIP_DEFLATED
from docx import Document
from pptx import Presentation
from pptx.util import Inches
from openpyxl import Workbook

root = Path(__file__).resolve().parent
names = ['source.docx', 'source.pptx', 'source.xlsx', 'source.odt', 'source.rtf']
if any((root / name).exists() for name in names):
    raise SystemExit('Refusing to overwrite existing original fixtures')
doc = Document()
doc.add_heading('Original document', 0)
doc.add_paragraph('Original body A&B -- spaces. 中文。')
doc.add_table(rows=2, cols=2).cell(0, 0).text = 'Original table'
doc.core_properties.created = datetime(2002, 3, 4, tzinfo=timezone.utc)
doc.save(root / names[0])
pres = Presentation()
slide = pres.slides.add_slide(pres.slide_layouts[5])
slide.shapes.title.text = 'Original presentation'
slide.shapes.add_textbox(Inches(1), Inches(2), Inches(6), Inches(1)).text = 'Original slide A&B -- 中文。'
pres.save(root / names[1])
book = Workbook()
book.active['A1'] = 'Original workbook A&B -- 中文。'
book.active['B1'] = 42
book.active['B2'] = '=B1*2'
book.save(root / names[2])
with ZipFile(root / names[3], 'x') as archive:
    archive.writestr('mimetype', 'application/vnd.oasis.opendocument.text', compress_type=ZIP_STORED)
    archive.writestr('META-INF/manifest.xml', '''<manifest:manifest xmlns:manifest="urn:oasis:names:tc:opendocument:xmlns:manifest:1.0" manifest:version="1.2"><manifest:file-entry manifest:full-path="/" manifest:media-type="application/vnd.oasis.opendocument.text"/><manifest:file-entry manifest:full-path="content.xml" manifest:media-type="text/xml"/><manifest:file-entry manifest:full-path="meta.xml" manifest:media-type="text/xml"/></manifest:manifest>''', compress_type=ZIP_DEFLATED)
    archive.writestr('meta.xml', '''<office:document-meta xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" office:version="1.2"><office:meta/></office:document-meta>''', compress_type=ZIP_DEFLATED)
    archive.writestr('content.xml', '''<office:document-content xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" office:version="1.2"><office:body><office:text><text:p>Original ODT A&amp;B -- 中文。</text:p></office:text></office:body></office:document-content>''', compress_type=ZIP_DEFLATED)
(root / names[4]).write_bytes(br"{\rtf1\ansi{\fonttbl{\f0 Arial;}}\f0 Original body A\&B -- \'e9.}")
